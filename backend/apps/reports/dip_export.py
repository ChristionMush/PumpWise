"""Dip history as an Excel workbook (managers and owners), to check and inspect every dip.

Sheets:
* Dips: every dip in time order per tank, with two checks:
  - chart check: the litres the tank's calibration chart gives for the mm entered, against the litres recorded (a typo
    such as 1928 for 2928 stands out);
  - stock check, from one day-close dip to the next: previous dip + deliveries - litres sold by the close meters =
    expected litres, against the litres dipped (minus = less fuel in the tank than expected).
* Deliveries, Corrections (changes made to a day close), Days without a dip, and Tanks (the last dip of each).

Built from the same records as the Dip history screen. Read-only.
"""
import io
from datetime import datetime, time, timedelta
from decimal import Decimal
from types import SimpleNamespace

from django.utils import timezone
from openpyxl import Workbook
from rest_framework.exceptions import ValidationError

from apps.operations.models import DayCloseReading, Delivery, TankReading
from apps.stations.models import Tank
from apps.stations.services import litres_from_dip

from .accounting import GAIN_FILL, LOSS_FILL, _local
from .accounting import _sheet as _pack_sheet
from .tank_history import tank_history

DAY = "ddd dd mmm yyyy"
CHART_TOLERANCE = Decimal("20")  # litres: a recorded figure further than this from the chart is worth a second look


def _sheet(wb, title, station, period, columns, rows, fills=None, formats=None):
    ws = _pack_sheet(wb, title, station, period, columns, rows, fills)
    if not rows:
        ws["A4"] = "Nothing recorded in this period."
    for col, fmt in (formats or {}).items():  # dates as dates, not "00:00"
        for (cell,) in ws.iter_rows(min_row=4, min_col=col, max_col=col):
            cell.number_format = fmt
    return ws


def _f(v, places=3):
    return float(Decimal(v).quantize(Decimal(10) ** -places)) if v is not None else None


def _chart(tank, mm):
    if mm is None:
        return None
    try:
        return litres_from_dip(tank, mm)
    except ValidationError:
        return None  # no chart, or the mm is outside it


def _close_meters(close_id, tank):
    rows = DayCloseReading.objects.filter(day_close_id=close_id, nozzle__tank=tank)
    return {r.nozzle_id: r.reading for r in rows}


def _dip_rows(tanks, start, end):
    rows, fills = [], []
    for t in tanks:
        has_chart = t.calibration.count() >= 2
        before = TankReading.objects.filter(tank=t, reading_at__lt=start).order_by("-reading_at").first()
        dips = list(TankReading.objects.filter(tank=t, reading_at__gte=start, reading_at__lt=end).select_related("recorded_by").order_by("reading_at"))
        prev = before
        for d in dips:
            notes, fill = [], None
            chart = _chart(t, d.dip_mm)
            chart_diff = d.litres - chart if chart is not None else None
            if d.dip_mm is None:
                notes.append("no mm entered")
            elif chart is None and has_chart:
                notes.append("mm outside the calibration chart")
            elif abs(chart_diff) > CHART_TOLERANCE:
                notes.append("litres do not match the mm on the chart")
                fill = LOSS_FILL
            delivered = sold = expected = diff = None
            if prev is not None:
                delivered = sum((v.litres_received for v in Delivery.objects.filter(tank=t, delivered_at__gt=prev.reading_at, delivered_at__lte=d.reading_at)), Decimal(0))
                if d.day_close_id and prev.day_close_id:
                    now_m, then_m = _close_meters(d.day_close_id, t), _close_meters(prev.day_close_id, t)
                    if now_m and set(now_m) == set(then_m):
                        sold = sum((now_m[n] - then_m[n] for n in now_m), Decimal(0))
                        expected = prev.litres + delivered - sold
                        diff = d.litres - expected
                        if fill is None and diff:
                            fill = LOSS_FILL if diff < 0 else GAIN_FILL
                else:
                    notes.append("stock check only runs from one day-close dip to the next")
            local = timezone.localtime(d.reading_at)
            rows.append([
                local.date(), local.strftime("%H:%M"), t.station.name, t.name, t.fuel_type.code,
                _f(d.dip_mm, 1), _f(d.litres), _f(chart) if has_chart else "no chart", _f(chart_diff), "Yes" if d.day_close_id else "No",
                d.recorded_by.full_name if d.recorded_by_id else "",
                _f(prev.litres) if prev else None, _f(delivered), _f(sold, 2), _f(expected), _f(diff), "; ".join(notes),
            ])
            fills.append(fill)
            prev = d
    order = sorted(range(len(rows)), key=lambda i: (rows[i][2], rows[i][3], rows[i][0], rows[i][1]))
    return [rows[i] for i in order], [fills[i] for i in order]


def build_dip_export(*, station_ids, date_from, date_to):
    tanks = Tank.objects.select_related("station", "fuel_type").order_by("station__name", "name")
    if station_ids is not None:
        tanks = tanks.filter(station_id__in=station_ids)
    tanks = list(tanks)
    tz = timezone.get_current_timezone()
    start = timezone.make_aware(datetime.combine(date_from, time.min), tz)
    end = timezone.make_aware(datetime.combine(date_to + timedelta(days=1), time.min), tz)
    names = sorted({t.station.name for t in tanks})
    title = SimpleNamespace(name=names[0] if len(names) == 1 else "All stations")
    period = f"{date_from:%d %b %Y} to {date_to:%d %b %Y}"
    history = tank_history(station_ids=station_ids, date_from=date_from, date_to=date_to)

    wb = Workbook()
    wb.remove(wb.active)
    rows, fills = _dip_rows(tanks, start, end)
    _sheet(wb, "Dips", title, period, [
        "Date", "Time", "Station", "Tank", "Fuel", "Dip (mm)", "Litres recorded", "Litres by chart for the mm",
        "Chart check (recorded - chart, L)", "Day close dip", "Recorded by", "Previous dip (L)", "Delivered since (L)",
        "Sold by close meters since (L)", "Expected (L)", "Difference (L, - less than expected)", "Note"], rows, fills, {1: DAY})

    deliveries = []
    for v in Delivery.objects.filter(tank__in=tanks, delivered_at__gte=start, delivered_at__lt=end).select_related("tank__station", "recorded_by").order_by("delivered_at"):
        local = timezone.localtime(v.delivered_at)
        deliveries.append([local.date(), local.strftime("%H:%M"), v.tank.station.name, v.tank.name, _f(v.litres_received), v.supplier, v.note_number,
                           v.recorded_by.full_name if v.recorded_by_id else ""])
    _sheet(wb, "Deliveries", title, period, ["Date", "Time", "Station", "Tank", "Litres", "Supplier", "Note number", "Recorded by"], deliveries, formats={1: DAY})

    def dip_text(x):
        return f"{x['dip_mm']} mm = {x['litres']} L" if x.get("dip_mm") else f"{x.get('litres')} L"

    corrections = []
    for day in reversed(history["days"]):
        for e in day["entries"]:
            if e["type"] != "correction":
                continue
            changes = [f"{x['tank']} {dip_text(x['from'])} -> {dip_text(x['to'])}" for x in e.get("dips", [])]
            changes += [f"{x['pump']} meter {x['from']} -> {x['to']}" for x in e.get("meters", [])]
            when = timezone.localtime(datetime.fromisoformat(e["at"]))
            corrections.append([when.date(), when.strftime("%H:%M"), e.get("station"), e.get("close_day"), "; ".join(changes), e.get("note", ""), e.get("by") or ""])
    _sheet(wb, "Corrections", title, period, ["Date", "Time", "Station", "Day close corrected", "What changed", "Note", "By"], corrections, formats={1: DAY})

    # today is not over: its dip comes at the day close, about 21:00
    today = timezone.localdate()
    missing = [[datetime.fromisoformat(d["date"]).date(), m] for d in reversed(history["days"]) if d["date"] != today.isoformat() for m in d["missing"]]
    _sheet(wb, "Days without a dip", title, period, ["Date", "Station · tank"], missing, [LOSS_FILL] * len(missing), {1: DAY})

    tanks_rows = [[t["station"], t["name"], t["fuel_code"], _local(datetime.fromisoformat(t["last_dip_at"])) if t["last_dip_at"] else None,
                   t["last_dip_by"] or "", t["hours_since_last_dip"], "Overdue" if t["overdue"] else "On time"] for t in history["tanks"]]
    _sheet(wb, "Tanks", title, f"each tank should be dipped every {history['every_hours']} hours",
           ["Station", "Tank", "Fuel", "Last dip", "By", "Hours since", "Status"], tanks_rows,
           [LOSS_FILL if r[-1] == "Overdue" else None for r in tanks_rows], {4: "ddd dd mmm yyyy hh:mm"})

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), title.name
