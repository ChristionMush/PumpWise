"""Dip history: when each tank was dipped, by whom, deliveries, the days without a dip, and overdue tanks."""
from datetime import datetime, timedelta
from decimal import Decimal

from django.utils import timezone

from apps.accounts.models import Role
from apps.core.db import db_context
from apps.operations import services
from apps.stations.models import Tank

from .helpers import client_for, dip, make_forecourt, two_tenants


def local(days_ago, hour):
    day = timezone.localdate() - timedelta(days=days_ago)
    return timezone.make_aware(datetime(day.year, day.month, day.day, hour, 30))


def test_managers_see_each_dip_by_day_the_missing_days_and_overdue_tanks():
    a, b = two_tenants()
    fc = make_forecourt(a["tenant"], name="Msasa Road")
    other = make_forecourt(b["tenant"], name="Elsewhere")
    with db_context(tenant_id=a["tenant"].id):
        Tank.objects.filter(pk=fc.tank.pk).update(created_at=timezone.now() - timedelta(days=10))
    mgr = a[Role.MANAGER]
    dip(a["tenant"], mgr, fc, 6000, local(2, 6))
    dip(a["tenant"], mgr, fc, 5200, local(2, 15))
    with db_context(tenant_id=a["tenant"].id):
        services.record_delivery(user=mgr, tank=fc.tank, litres=Decimal("8000"), delivered_at=local(2, 9), supplier="Zuva", note_number="Z-1")
    dip(b["tenant"], b[Role.MANAGER], other, 100, local(1, 6))  # another business: never shown

    res = client_for(mgr).get("/api/reports/dip-history/", {"date_from": str(timezone.localdate() - timedelta(days=2))})
    assert res.status_code == 200, res.content
    body = res.json()
    assert body["every_hours"] == 24
    (tank,) = body["tanks"]
    assert tank["station"] == "Msasa Road" and tank["last_dip_by"] == mgr.full_name and tank["overdue"] is True  # last dip 2 days ago
    days = {d["date"]: d for d in body["days"]}
    two_ago = days[str(timezone.localdate() - timedelta(days=2))]
    assert [(e["type"], e.get("litres")) for e in two_ago["entries"]] == [("dip", "6000"), ("delivery", "8000"), ("dip", "5200")]
    assert two_ago["missing"] == []
    assert days[str(timezone.localdate() - timedelta(days=1))]["missing"] == ["Msasa Road · Tank 1"]  # no dip yesterday
    assert list(days) == sorted(days, reverse=True)  # newest day first


def test_a_fresh_dip_is_not_overdue_and_attendants_cannot_see_dip_history():
    a, _ = two_tenants()
    fc = make_forecourt(a["tenant"])
    dip(a["tenant"], a[Role.MANAGER], fc, 5000, timezone.now() - timedelta(hours=3))
    (tank,) = client_for(a[Role.OWNER]).get("/api/reports/dip-history/").json()["tanks"]
    assert tank["overdue"] is False and 2.9 <= tank["hours_since_last_dip"] <= 3.1
    assert client_for(a[Role.ATTENDANT]).get("/api/reports/dip-history/").status_code == 403


def test_the_dip_sheet_lists_every_dip_with_a_chart_check():
    import io

    from openpyxl import load_workbook

    a, _ = two_tenants()
    fc = make_forecourt(a["tenant"], name="Msasa Road")
    mgr = a[Role.MANAGER]
    dip(a["tenant"], mgr, fc, 6000, local(2, 6))
    dip(a["tenant"], mgr, fc, 4000, local(1, 6))
    assert client_for(mgr).get("/api/reports/dip-history/export/").status_code == 403  # the owner's
    res = client_for(a[Role.OWNER]).get("/api/reports/dip-history/export/", {"date_from": str(timezone.localdate() - timedelta(days=2))})
    assert res.status_code == 200 and res["Content-Type"].startswith("application/vnd.openxml"), res.content[:200]
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Dips", "Deliveries", "Corrections", "Days without a dip", "Tanks"]
    rows = [r for r in wb["Dips"].iter_rows(min_row=4, values_only=True)]
    assert [(r[3], r[6], r[11]) for r in rows] == [("Tank 1", 6000, None), ("Tank 1", 4000, 6000)]  # tank, litres, previous dip
    assert rows[0][7] is None and rows[0][16] == "no mm entered"  # the chart check needs the mm
    assert client_for(a[Role.ATTENDANT]).get("/api/reports/dip-history/export/").status_code == 403



def test_the_dip_sheet_works_for_a_tank_without_a_calibration_chart():
    import io

    from openpyxl import load_workbook

    from apps.stations.models import CalibrationPoint

    a, _ = two_tenants()
    fc = make_forecourt(a["tenant"])
    with db_context(tenant_id=a["tenant"].id):
        CalibrationPoint.objects.filter(tank=fc.tank).delete()
        services.record_tank_dip(user=a[Role.MANAGER], tank=fc.tank, litres=Decimal("4000"), dip_mm=Decimal("800"), at=local(1, 6))
    res = client_for(a[Role.OWNER]).get("/api/reports/dip-history/export/", {"date_from": str(timezone.localdate() - timedelta(days=2))})
    assert res.status_code == 200, res.content[:300]
    (row,) = load_workbook(io.BytesIO(res.content))["Dips"].iter_rows(min_row=4, values_only=True)
    assert (row[5], row[6], row[7], row[16]) == (800, 4000, "no chart", None)  # no note: an empty cell
