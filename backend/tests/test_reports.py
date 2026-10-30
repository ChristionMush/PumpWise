"""Phase 4 (daily sales) and Phase 5 (reconciliation)."""
from decimal import Decimal

from apps.accounts.models import Role
from apps.core.db import db_context

from .helpers import at, client_for, dip, make_forecourt, run_shift, two_tenants


def test_daily_sales_group_by_local_day_pump_fuel_and_use_the_price_in_force():
    a, _ = two_tenants()
    fc = make_forecourt(a["tenant"], price="1.00", price_from=at(2026, 9, 1))
    owner = client_for(a[Role.OWNER])
    owner.post("/api/prices/", {"station": str(fc.station.pk), "fuel_type": str(fc.fuel.pk), "price_per_litre": "2.00", "effective_from": "2026-10-02T00:00:00Z"}, format="json")
    # 22:30 UTC on 30 Sep is 00:30 on 1 Oct in Harare (UTC+2): it belongs to 1 October
    run_shift(a["tenant"], a[Role.ATTENDANT], fc, [0, 0], [100, 40], started_at=at(2026, 9, 30, 14), ended_at=at(2026, 9, 30, 22, 30))
    run_shift(a["tenant"], a[Role.ATTENDANT], fc, [100, 40], [150, 40], started_at=at(2026, 10, 2, 6), ended_at=at(2026, 10, 2, 12))
    res = owner.get(f"/api/reports/daily-sales/?station={fc.station.pk}&date_from=2026-09-30&date_to=2026-10-02")
    assert res.status_code == 200, res.content
    days = {d["date"]: d for d in res.json()["days"]}
    assert days["2026-09-30"]["litres"] == "0.000"
    assert days["2026-10-01"]["litres"] == "140.000" and days["2026-10-01"]["amount_by_currency"] == {"USD": "140.00"}
    assert days["2026-10-02"]["litres"] == "50.000" and days["2026-10-02"]["amount_by_currency"] == {"USD": "100.00"}  # new price
    assert days["2026-10-01"]["by_pump"][0]["pump_name"] == "Pump 1" and days["2026-10-01"]["by_fuel"][0]["fuel_code"] == "PETROL"
    assert res.json()["totals"]["litres"] == "190.000"


def test_averages_per_day_and_per_weekday_include_quiet_days(monkeypatch):
    a, _ = two_tenants()
    fc = make_forecourt(a["tenant"], price="1.00", price_from=at(2026, 9, 1))
    # Thursdays 1 Oct and 8 Oct: 100 litres each; other days nothing
    run_shift(a["tenant"], a[Role.ATTENDANT], fc, [0, 0], [100, 0], started_at=at(2026, 10, 1, 6), ended_at=at(2026, 10, 1, 12))
    run_shift(a["tenant"], a[Role.ATTENDANT], fc, [100, 0], [200, 0], started_at=at(2026, 10, 8, 6), ended_at=at(2026, 10, 8, 12))
    from django.utils import timezone

    monkeypatch.setattr(timezone, "localdate", lambda *_a, **_k: __import__("datetime").date(2026, 10, 14))
    d = client_for(a[Role.OWNER]).get("/api/reports/daily-sales/?date_from=2026-10-01&date_to=2026-10-14").json()
    assert d["averages"]["avg_daily_litres"] == "14.286"  # 200 litres over 14 days
    wd = {w["name"]: w for w in d["averages"]["by_weekday"]}
    assert wd["Thursday"]["avg_litres"] == "100.000" and wd["Thursday"]["days"] == 2
    assert wd["Monday"]["avg_litres"] == "0.000" and wd["Monday"]["days"] == 2
    assert d["averages"]["by_fuel"]["PETROL"]["by_weekday"][3]["avg_litres"] == "100.000"


def test_reports_respect_roles_dates_and_tenants():
    a, b = two_tenants()
    fa = make_forecourt(a["tenant"])
    make_forecourt(b["tenant"])
    run_shift(a["tenant"], a[Role.ATTENDANT], fa, [0, 0], [50, 0], started_at=at(2026, 10, 1, 6), ended_at=at(2026, 10, 1, 12))
    url = "/api/reports/daily-sales/?date_from=2026-10-01&date_to=2026-10-01"
    assert client_for(a[Role.ATTENDANT]).get(url).status_code == 403
    assert client_for(a[Role.MANAGER]).get(url).status_code == 200
    assert client_for(a[Role.OWNER]).get(url.replace("date_to=2026-10-01", "date_to=bad")).status_code == 400
    assert client_for(a[Role.OWNER]).get("/api/reports/daily-sales/?date_from=2025-01-01&date_to=2026-10-01").status_code == 400  # > 1 year
    assert client_for(b[Role.OWNER]).get(url).json()["totals"]["litres"] == "0.000"  # B never sees A's sales
    assert client_for(b[Role.OWNER]).get(f"{url}&station={fa.station.pk}").status_code == 200  # unknown to B: simply nothing
    assert client_for(b[Role.OWNER]).get(f"{url}&station={fa.station.pk}").json()["totals"]["litres"] == "0.000"


def _day_with_sales(a, fc, sold, *, opening, closing, delivered=0, day=1):
    owner = a[Role.OWNER]
    dip(a["tenant"], owner, fc, opening, at(2026, 10, day, 5, 0))
    run_shift(a["tenant"], a[Role.ATTENDANT], fc, [0, 0], [sold, 0], started_at=at(2026, 10, day, 6), ended_at=at(2026, 10, day, 12))
    if delivered:
        from apps.operations import services

        with db_context(tenant_id=a["tenant"].id):
            services.record_delivery(user=owner, tank=fc.tank, litres=delivered, delivered_at=at(2026, 10, day, 8))
    dip(a["tenant"], owner, fc, closing, at(2026, 10, day, 18, 0))  # 20:00 Harare, before the day closes at 21:00


def test_tank_loss_is_opening_plus_deliveries_minus_sold_minus_closing():
    a, _ = two_tenants()
    fc = make_forecourt(a["tenant"])
    _day_with_sales(a, fc, 1000, opening=6000, closing=8000, delivered=3500)  # 6000 + 3500 - 1000 - 8000 = 500 litres missing
    rows = client_for(a[Role.OWNER]).get("/api/reconciliation/tanks/?date_from=2026-10-01&date_to=2026-10-01").json()["rows"]
    r = [x for x in rows if x["date"] == "2026-10-01"][0]
    assert (r["opening"], r["deliveries"], r["dispensed"], r["closing"], r["loss"]) == ("6000.000", "3500.000", "1000.000", "8000.000", "500.000")
    assert r["status"] == "loss" and r["flagged"] is True and r["loss_pct"] == "50.00"


def test_small_differences_are_ok_and_the_threshold_is_configurable():
    a, _ = two_tenants()
    fc = make_forecourt(a["tenant"])
    _day_with_sales(a, fc, 1000, opening=6000, closing=4985)  # 15 litres unexplained: under the default 20 L
    owner = client_for(a[Role.OWNER])
    row = lambda: owner.get("/api/reconciliation/tanks/?date_from=2026-10-01&date_to=2026-10-01").json()["rows"][0]  # noqa: E731
    assert row()["status"] == "ok" and row()["loss"] == "15.000"
    from apps.tenants.services import get_settings

    with db_context(tenant_id=a["tenant"].id):
        s = get_settings(a["tenant"].id)
        s.loss_threshold_litres, s.loss_threshold_pct = Decimal("5"), Decimal("0.1")
        s.save()
    assert row()["status"] == "loss" and row()["flagged"] is True
    assert owner.get("/api/reconciliation/tanks/?date_from=2026-10-01&date_to=2026-10-01&flagged=1").json()["rows"]


def test_a_gain_is_flagged_too_because_it_means_a_wrong_number_somewhere():
    a, _ = two_tenants()
    fc = make_forecourt(a["tenant"])
    _day_with_sales(a, fc, 100, opening=5000, closing=5000)  # 100 sold but the tank did not drop
    r = client_for(a[Role.OWNER]).get("/api/reconciliation/tanks/?date_from=2026-10-01&date_to=2026-10-01").json()["rows"][0]
    assert r["status"] == "gain" and r["flagged"] and r["loss"] == "-100.000"


def test_not_enough_dips_means_no_verdict_rather_than_a_guess():
    a, _ = two_tenants()
    fc = make_forecourt(a["tenant"])
    run_shift(a["tenant"], a[Role.ATTENDANT], fc, [0, 0], [100, 0], started_at=at(2026, 10, 1, 6), ended_at=at(2026, 10, 1, 12))
    assert client_for(a[Role.OWNER]).get("/api/reconciliation/tanks/?date_from=2026-10-01&date_to=2026-10-01").json()["rows"] == []


def test_cash_reconciliation_lists_closed_shifts_and_filters_flagged():
    a, _ = two_tenants()
    fc = make_forecourt(a["tenant"], price="2.00")
    run_shift(a["tenant"], a[Role.ATTENDANT], fc, [0, 0], [100, 0], started_at=at(2026, 10, 1, 6), ended_at=at(2026, 10, 1, 12), cash=[("cash", "USD", 200)])
    run_shift(a["tenant"], a[Role.ATTENDANT], fc, [100, 0], [200, 0], started_at=at(2026, 10, 2, 6), ended_at=at(2026, 10, 2, 12), cash=[("cash", "USD", 150)])
    owner = client_for(a[Role.OWNER])
    rows = owner.get("/api/reconciliation/cash/?date_from=2026-10-01&date_to=2026-10-02").json()["rows"]
    assert {r["status"] for r in rows} == {"ok", "short"}
    flagged = owner.get("/api/reconciliation/cash/?date_from=2026-10-01&date_to=2026-10-02&flagged=1").json()["rows"]
    assert len(flagged) == 1 and flagged[0]["variance"] == "-50.00" and flagged[0]["attendant"]
    assert client_for(a[Role.ATTENDANT]).get("/api/reconciliation/cash/").status_code == 403


def test_the_monthly_accounting_pack_has_every_sheet_with_the_months_numbers():
    import io

    from openpyxl import load_workbook

    from apps.operations import services

    a, _ = two_tenants()
    fc = make_forecourt(a["tenant"], price="1.50")
    owner = client_for(a[Role.OWNER])
    dip(a["tenant"], a[Role.OWNER], fc, 8000, at(2026, 10, 1, 4))
    run_shift(a["tenant"], a[Role.ATTENDANT], fc, [0, 0], [100, 0], started_at=at(2026, 10, 1, 5), ended_at=at(2026, 10, 1, 15), cash=[("cash", "USD", 140)])
    with db_context(tenant_id=a["tenant"].id):
        services.record_delivery(user=a[Role.OWNER], tank=fc.tank, litres=Decimal(1000), delivered_at=at(2026, 10, 1, 16), supplier="Zuva", note_number="DN-7")
    dip(a["tenant"], a[Role.OWNER], fc, 8900, at(2026, 10, 1, 20))  # 8000 - 100 sold + 1000 delivered: balances

    res = owner.get("/api/reports/accounting-pack/", {"station": str(fc.station.pk), "month": "2026-10"})
    assert res.status_code == 200 and "spreadsheetml" in res["Content-Type"] and "2026-10" in res["Content-Disposition"]
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Daily", "Sales", "Takings", "Cash differences", "Deliveries", "Stock", "Prices", "Notes"]

    def first(name):
        return next(wb[name].iter_rows(min_row=4, values_only=True))

    # the trading day: litres, sales, grand total by payment type, cash losses ($10 short is the attendant's), the dip
    # Grand Total = the fuel sold (150): 140 handed in + 10 short
    assert first("Daily")[1:] == ("not closed", None, 100.0, 150.0, 150.0, 140.0, 0.0, 0.0, 0.0, 0.0, 10.0, 10.0, 10.0, 0.0, 8000.0)
    assert first("Sales")[1:] == ("Petrol", 100.0, 1.5, 150.0, "USD")
    assert first("Takings")[1:] == ("Cash", "USD", 140.0, 1.0, 140.0)
    assert first("Cash differences")[4:] == (150.0, 140.0, -10.0, "Short")
    assert first("Deliveries")[1:6] == ("Tank 1", "Petrol", "Zuva", "DN-7", 1000.0)
    assert first("Stock")[1:] == ("Tank 1", "Petrol", 8000.0, 1000.0, 100.0, 8900.0, 0.0, "OK")
    assert first("Prices")[:3] == ("Petrol", "USD", 1.5)
    assert first("Notes")[0] == "Nothing recorded in this month."

    for other in (a[Role.MANAGER], a[Role.ATTENDANT]):  # the owner's only
        assert client_for(other).get("/api/reports/accounting-pack/", {"station": str(fc.station.pk), "month": "2026-10"}).status_code == 403
    assert owner.get("/api/reports/accounting-pack/", {"station": str(fc.station.pk), "month": "October"}).status_code == 400
    assert owner.get("/api/reports/accounting-pack/", {"month": "2026-10"}).status_code == 400


def test_an_owners_note_for_the_books_shows_in_the_accounting_pack_and_the_day_log():
    import io
    from datetime import date

    from openpyxl import load_workbook

    from apps.reports.notes import add_book_note

    a, _ = two_tenants()
    fc = make_forecourt(a["tenant"], price="2.05")
    owner = client_for(a[Role.OWNER])
    note = "Pump 4A, night of 28-29 Oct: 224.15 L sold were not recorded under the attendant who sold them."
    with db_context(tenant_id=a["tenant"].id):
        add_book_note(station=fc.station, day=date(2026, 10, 29), note=note, litres=Decimal("224.15"), amount=Decimal("459.51"), actor=a[Role.OWNER])
        add_book_note(station=fc.station, day=date(2026, 11, 2), note="next month", actor=a[Role.OWNER])

    res = owner.get("/api/reports/accounting-pack/", {"station": str(fc.station.pk), "month": "2026-10"})
    rows = list(load_workbook(io.BytesIO(res.content))["Notes"].iter_rows(min_row=4, values_only=True))
    assert len(rows) == 1 and rows[0][1:6] == (note, 224.15, 459.51, "USD", a[Role.OWNER].full_name)
    assert rows[0][0].date() == date(2026, 10, 29)

    log = owner.get("/api/reports/day-log/", {"station": str(fc.station.pk), "day": "2026-10-29"}).json()
    assert f"Note for the books: {note}" in log["highlights"]
    other = owner.get("/api/reports/day-log/", {"station": str(fc.station.pk), "day": "2026-10-28"}).json()
    assert not any("Note for the books" in h for h in other["highlights"])


def test_a_delivery_must_go_into_a_tank_of_the_fuel_received():
    a, _ = two_tenants()
    fc = make_forecourt(a["tenant"])
    owner = client_for(a[Role.OWNER])
    body = {"tank": str(fc.tank.pk), "litres": "1000", "supplier": "Zuva"}
    wrong = owner.post("/api/deliveries/", {**body, "fuel_code": "DIESEL"}, format="json")
    assert wrong.status_code == 400 and "not DIESEL" in str(wrong.json())
    assert owner.post("/api/deliveries/", {**body, "fuel_code": "PETROL"}, format="json").status_code == 201


def test_intercompany_issues_are_owed_to_the_business_not_money_received():
    import io

    from openpyxl import load_workbook

    a, _ = two_tenants()
    fc = make_forecourt(a["tenant"], price="1.50")
    run_shift(a["tenant"], a[Role.ATTENDANT], fc, [0, 0], [100, 0], started_at=at(2026, 10, 1, 5), ended_at=at(2026, 10, 1, 15),
              cash=[("cash", "USD", 100), ("credit", "USD", 50)])  # 150 expected: 100 in cash, 50 issued to another company
    res = client_for(a[Role.OWNER]).get("/api/reports/accounting-pack/", {"station": str(fc.station.pk), "month": "2026-10"})
    rows = {r[0]: r[5] for r in load_workbook(io.BytesIO(res.content))["Takings"].iter_rows(min_row=4, values_only=True) if isinstance(r[0], str)}
    assert rows["Money received"] == 100.0 and rows["Total handed in"] == 150.0
    assert [v for k, v in rows.items() if k.startswith("Owed to")] == [50.0]
